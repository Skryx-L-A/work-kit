import io
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from kitdesign import html_deck, pptx_build
from kitdesign.cli import main
from kitdesign.deck import parse
from kitdesign.refs import load_theme

DECK = "# Title\n\nsub\n\n---\n# Point\n\n- a\n- b\n\n::: notes\nnote text\n:::\n\n---\n# Img\n\n![c](missing.png)\n"

PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cfc0f01f0005000201a0d1d5b10000000049454e44ae426082"
)


def test_clean_pptx(tmp_path):
    out = tmp_path / "d.pptx"
    warnings = pptx_build.build(parse(DECK, tmp_path), out, load_theme())
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    assert prs.slide_width == Inches(13.333)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "note text"
    assert any("missing.png" in w for w in warnings)


def _template(tmp_path, potx=False):
    base = tmp_path / "base.pptx"
    Presentation().save(str(base))
    if not potx:
        return base
    dst = tmp_path / "company.potx"
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(dst, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            z.writestr(item, data)
    return dst


def test_template_pptx_and_potx(tmp_path):
    for potx in (False, True):
        out = tmp_path / f"t{potx}.pptx"
        pptx_build.build(parse(DECK, tmp_path), out, load_theme(), _template(tmp_path, potx))
        prs = Presentation(str(out))
        assert len(prs.slides) == 3
        texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
        assert "Point" in texts and any("a" in t and "b" in t for t in texts)


def test_html_is_self_contained(tmp_path):
    (tmp_path / "p.png").write_bytes(PNG)
    deck = parse("# T\n\n---\n# Pic\n\n![x](p.png)\n", tmp_path)
    out = tmp_path / "d.html"
    assert html_deck.build(deck, out, load_theme()) == []
    page = out.read_text()
    assert page.count('<section class="slide') == 2
    assert "data:image/png;base64," in page
    assert not re.search(r'(src|href)="https?:', page)


def test_theme_json(refs_home):
    (refs_home / "brand" / "theme.json").write_text('{"colors": {"accent": "#AA0000", "text": "red"}, "fonts": {"body": "Inter"}}')
    t = load_theme()
    assert t.accent == "#aa0000" and t.body_font == "Inter"
    assert any("colors.text" in w for w in t.warnings)
    assert '--accent: #aa0000' in t.css_vars()


def test_cli_auto_template(tmp_path, refs_home, capsys):
    src = tmp_path / "deck.md"
    src.write_text(DECK)
    tpl = _template(tmp_path)
    (refs_home / "slides" / "company.pptx").write_bytes(tpl.read_bytes())
    assert main(["deck", str(src), "-o", str(tmp_path / "o.pptx")]) == 0
    assert "using template from design-refs" in capsys.readouterr().err
    (refs_home / "slides" / "other.pptx").write_bytes(tpl.read_bytes())
    assert main(["deck", str(src), "-o", str(tmp_path / "o.pptx")]) == 0
    assert "several templates" in capsys.readouterr().err


def test_cli_refs_and_new(tmp_path, refs_home, capsys):
    assert main(["refs", "--json"]) == 0
    assert '"slides": []' in capsys.readouterr().out
    (refs_home / "web" / "shot.png").write_bytes(PNG)
    (refs_home / "brand" / "theme.json").write_text('{"colors": {"accent": "#123456"}}')
    assert main(["new", "web", str(tmp_path / "w")]) == 0
    out = capsys.readouterr().out
    assert "design references found" in out
    assert "#123456" in (tmp_path / "w" / "tokens.css").read_text()
    assert main(["new", "web", str(tmp_path / "w")]) == 1  # no silent overwrite


def test_doc_html(tmp_path):
    from kitdesign.doc import to_html
    src = tmp_path / "r.md"
    src.write_text('---\ntitle: "R: one"\nversion: v2\n---\n\n## A\n\n| x | y |\n|---|--:|\n| 1 | 2 |\n')
    page, warnings = to_html(src, load_theme())
    assert "<table>" in page and "<title>R: one</title>" in page
    assert "/*FOOTER*/" not in page and "R: one · v2" in page
    assert warnings == []


def test_doc_heading_is_used_as_title_without_front_matter(tmp_path):
    from kitdesign.doc import to_html
    src = tmp_path / "bericht.md"
    src.write_text("# Wochenbericht\n\nStand der Arbeit.\n")
    page, warnings = to_html(src, load_theme())
    assert "<title>Wochenbericht</title>" in page
    assert '<header class="doc-head"><h1>Wochenbericht</h1>' in page
    assert "<title>bericht</title>" not in page
    assert page.count("Wochenbericht</h1>") == 1
    assert "Stand der Arbeit." in page
    assert warnings == []
